import os
import csv
import json
import re
import math
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, date
from django.utils import timezone
import pytz
from openpyxl import load_workbook
from decimal import Decimal
from django.db import transaction
from django.contrib.auth.models import User
import uuid

from .models import Tenant, RawPayload, Airport, EmissionRecord, AuditLog, IngestionBatch
from .column_resolver import resolve_columns, fuzzy_resolve, _ALIAS_TO_CANONICAL
from .date_parser import parse_date as defra_parse_date
from .unit_normalizer import (
    normalize_fuel, normalize_electricity,
    normalize_flight, normalize_hotel, normalize_ground,
    infer_fuel_type,
)

# Static embedded database of major global airports to auto-seed coordinates
STATIC_AIRPORTS = [
    {"iata_code": "DEL", "latitude": 28.5665, "longitude": 77.1031, "city": "Delhi", "country": "IN"},
    {"iata_code": "BOM", "latitude": 19.0896, "longitude": 72.8656, "city": "Mumbai", "country": "IN"},
    {"iata_code": "BLR", "latitude": 13.1986, "longitude": 77.7066, "city": "Bengaluru", "country": "IN"},
    {"iata_code": "MAA", "latitude": 12.9941, "longitude": 80.1709, "city": "Chennai", "country": "IN"},
    {"iata_code": "HYD", "latitude": 17.2403, "longitude": 78.4294, "city": "Hyderabad", "country": "IN"},
    {"iata_code": "PNQ", "latitude": 18.5793, "longitude": 73.9089, "city": "Pune", "country": "IN"},
    {"iata_code": "CCU", "latitude": 22.6547, "longitude": 88.4467, "city": "Kolkata", "country": "IN"},
    {"iata_code": "AMD", "latitude": 23.0770, "longitude": 72.6347, "city": "Ahmedabad", "country": "IN"},
    {"iata_code": "LHR", "latitude": 51.4775, "longitude": -0.4614, "city": "London Heathrow", "country": "GB"},
    {"iata_code": "LGW", "latitude": 51.1537, "longitude": -0.1821, "city": "London Gatwick", "country": "GB"},
    {"iata_code": "MAN", "latitude": 53.3537, "longitude": -2.2750, "city": "Manchester", "country": "GB"},
    {"iata_code": "SFO", "latitude": 37.6213, "longitude": -122.3790, "city": "San Francisco", "country": "US"},
    {"iata_code": "JFK", "latitude": 40.6413, "longitude": -73.7781, "city": "New York JFK", "country": "US"},
    {"iata_code": "LAX", "latitude": 33.9425, "longitude": -118.4081, "city": "Los Angeles", "country": "US"},
    {"iata_code": "ORD", "latitude": 41.9742, "longitude": -87.9073, "city": "Chicago", "country": "US"},
    {"iata_code": "SIN", "latitude": 1.3644, "longitude": 103.9915, "city": "Singapore Changi", "country": "SG"},
    {"iata_code": "CDG", "latitude": 49.0097, "longitude": 2.5479, "city": "Paris CDG", "country": "FR"},
    {"iata_code": "AMS", "latitude": 52.3086, "longitude": 4.7639, "city": "Amsterdam", "country": "NL"},
    {"iata_code": "FRA", "latitude": 50.0379, "longitude": 8.5622, "city": "Frankfurt", "country": "DE"},
    {"iata_code": "MUC", "latitude": 48.3538, "longitude": 11.7861, "city": "Munich", "country": "DE"},
    {"iata_code": "ZRH", "latitude": 47.4647, "longitude": 8.5492, "city": "Zurich", "country": "CH"},
    {"iata_code": "DXB", "latitude": 25.2532, "longitude": 55.3657, "city": "Dubai", "country": "AE"},
    {"iata_code": "DOH", "latitude": 25.2609, "longitude": 51.6138, "city": "Doha", "country": "QA"},
    {"iata_code": "HKG", "latitude": 22.3080, "longitude": 113.9185, "city": "Hong Kong", "country": "HK"},
    {"iata_code": "NRT", "latitude": 35.7720, "longitude": 140.3929, "city": "Tokyo Narita", "country": "JP"},
    {"iata_code": "BKK", "latitude": 13.6811, "longitude": 100.7472, "city": "Bangkok", "country": "TH"},
    {"iata_code": "YYZ", "latitude": 43.6777, "longitude": -79.6248, "city": "Toronto", "country": "CA"},
    {"iata_code": "GRU", "latitude": -23.4356, "longitude": -46.4731, "city": "Sao Paulo", "country": "BR"},
]

def auto_seed_airports():
    """Dynamically seeds any missing airports to guarantee haversine resolutions."""
    for ap in STATIC_AIRPORTS:
        Airport.objects.get_or_create(
            iata_code=ap["iata_code"],
            defaults={
                "latitude": Decimal(str(ap["latitude"])),
                "longitude": Decimal(str(ap["longitude"])),
                "city": ap["city"],
                "country": ap["country"]
            }
        )

class ColumnNormalizer:
    @classmethod
    def resolve_header(cls, header: str) -> str:
        if not header:
            return ""
        canonical = _ALIAS_TO_CANONICAL.get(str(header).lower())
        if not canonical:
            canonical = fuzzy_resolve(str(header))
        return canonical or str(header).strip().lower().replace(".", "_").replace(" ", "_").replace("\"", "")

def sanitize_decimal(val_str: str) -> Decimal:
    if not val_str:
        return Decimal('0.0000')
    cleaned = str(val_str).strip().replace('"', '').replace(' ', '')
    if not cleaned:
        return Decimal('0.0000')
    
    if ',' in cleaned and '.' in cleaned:
        if cleaned.rfind(',') > cleaned.rfind('.'):
            cleaned = cleaned.replace('.', '').replace(',', '.')
        else:
            cleaned = cleaned.replace(',', '')
    elif ',' in cleaned:
        parts = cleaned.split(',')
        if len(parts[-1]) == 2:
            cleaned = cleaned.replace(',', '.')
        else:
            cleaned = cleaned.replace(',', '')
            
    try:
        return Decimal(cleaned)
    except Exception:
        return Decimal('0.0000')

def ensure_text_content(content: str) -> str:
    if not content:
        return ""
    content_stripped = content.strip()
    if content_stripped.startswith(("{", "[", "<")):
        return content
    
    # Try to decode base64 in case it is base64-encoded binary payload
    import base64
    try:
        decoded_bytes = base64.b64decode(content_stripped, validate=True)
        return decoded_bytes.decode('utf-8', errors='ignore')
    except Exception:
        return content

def parse_microsoft_epoch(date_str) -> datetime:
    if not date_str:
        return timezone.now()
    
    date_str = str(date_str).strip()
    if '/Date(' in date_str:
        match = re.search(r'/Date\((\d+)\)/', date_str)
        if match:
            epoch_ms = int(match.group(1))
            return datetime.fromtimestamp(epoch_ms / 1000.0, tz=pytz.utc)
            
    try:
        return parse_date_flexible(date_str)
    except Exception:
        return timezone.now()

def parse_date_flexible(date_str: str) -> datetime:
    dt_val, flag = defra_parse_date(date_str)
    if not dt_val:
        return timezone.now()
    if isinstance(dt_val, date) and not isinstance(dt_val, datetime):
        dt_val = datetime.combine(dt_val, datetime.min.time())
    if timezone.is_naive(dt_val):
        return timezone.make_aware(dt_val, pytz.utc)
    return dt_val

def haversine_distance(lat1, lon1, lat2, lon2):
    R = 6371.0
    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_theta = math.radians(lon2 - lon1)
    
    a = math.sin(delta_phi / 2.0)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(delta_theta / 2.0)**2
    c = 2 * math.asin(math.sqrt(a))
    return R * c

class BulkIngestionContext:
    def __init__(self):
        self.records = []
        self.audit_logs = []
        self.orig_record_create = None
        self.orig_audit_create = None

    def __enter__(self):
        self.orig_record_create = EmissionRecord.objects.create
        self.orig_audit_create = AuditLog.objects.create

        def custom_record_create(*args, **kwargs):
            obj = EmissionRecord(*args, **kwargs)
            if not obj.id:
                obj.id = uuid.uuid4()
            self.records.append(obj)
            return obj

        def custom_audit_create(*args, **kwargs):
            obj = AuditLog(*args, **kwargs)
            self.audit_logs.append(obj)
            return obj

        EmissionRecord.objects.create = custom_record_create
        AuditLog.objects.create = custom_audit_create
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        EmissionRecord.objects.create = self.orig_record_create
        AuditLog.objects.create = self.orig_audit_create

        if exc_type is None:
            if self.records:
                EmissionRecord.objects.bulk_create(self.records)
            if self.audit_logs:
                AuditLog.objects.bulk_create(self.audit_logs)

def parse_payload(raw_payload_id) -> dict:
    raw = RawPayload.objects.get(id=raw_payload_id)
    content = ensure_text_content(raw.payload_content)
    source = raw.source_system
    tenant = raw.tenant

    # Format auto-detection inside parser for resilience and database self-healing:
    content_stripped = content.strip()
    filename_lower = (raw.filename or "").lower()
    if content_stripped.startswith(("{", "[")):
        if "navanTmcResponse" in content_stripped or "navan" in filename_lower:
            source = RawPayload.SourceSystem.NAVAN_JSON
        else:
            source = RawPayload.SourceSystem.CONCUR_JSON
    elif content_stripped.startswith("<"):
        source = RawPayload.SourceSystem.SAP_IDOC
    elif content_stripped.startswith("EDI_DC40"):
        source = RawPayload.SourceSystem.SAP_IDOC
    elif filename_lower.endswith(".xlsx"):
        source = RawPayload.SourceSystem.SAP_XLSX
    else:
        # Sniff flat text files / CSV headers
        lines = content_stripped.splitlines()
        if lines:
            first_line_lower = lines[0].lower()
            if "mpan" in first_line_lower or "consumption_kwh" in first_line_lower or "standing_charge_gbp" in first_line_lower:
                source = RawPayload.SourceSystem.UTILITY_UK_CSV
            elif "tariff_band" in first_line_lower or "reading_from" in first_line_lower or "reading_to" in first_line_lower:
                source = RawPayload.SourceSystem.UTILITY_IN_CSV
            elif "trip_id" in first_line_lower and "expense_type" in first_line_lower:
                source = RawPayload.SourceSystem.TRAVEL_CSV
            elif "bukrs" in first_line_lower and "werks" in first_line_lower and "matnr" in first_line_lower:
                source = RawPayload.SourceSystem.SAP_CSV

    if raw.source_system != source:
        raw.source_system = source
        raw.save(update_fields=['source_system'])

    records_created = 0
    errors = []

    # Dynamic seed of expanded airport coordinate database
    auto_seed_airports()
    
    with transaction.atomic(), BulkIngestionContext() as bulk:
        locked_exist = EmissionRecord.objects.filter(raw_payload=raw, is_locked=True).exists()
        if locked_exist:
            raise exceptions.ValidationError("Cannot re-parse a raw payload that has already generated locked audit records.")
        
        EmissionRecord.objects.filter(raw_payload=raw).delete()

        # Resolve or Create IngestionBatch link
        batch_type_map = {
            RawPayload.SourceSystem.SAP_CSV: "SAP",
            RawPayload.SourceSystem.SAP_XLSX: "SAP",
            RawPayload.SourceSystem.SAP_ODATA: "SAP",
            RawPayload.SourceSystem.SAP_IDOC: "SAP",
            RawPayload.SourceSystem.UTILITY_IN_CSV: "UTILITY",
            RawPayload.SourceSystem.UTILITY_UK_CSV: "UTILITY",
            RawPayload.SourceSystem.TRAVEL_CSV: "TRAVEL",
            RawPayload.SourceSystem.CONCUR_JSON: "TRAVEL",
            RawPayload.SourceSystem.NAVAN_JSON: "TRAVEL",
        }
        source_type = batch_type_map.get(source, "SAP")

        # Resolve IngestionBatch linked to this raw payload
        batch, created = IngestionBatch.objects.get_or_create(
            raw_payload=raw,
            defaults={
                "tenant": tenant,
                "source_type": source_type,
                "uploaded_by": raw.ingested_by,
                "file_name": raw.filename or "API Bulk Feed",
                "status": "SUCCESS"
            }
        )

        if source == RawPayload.SourceSystem.SAP_CSV:
            records_created, errors = parse_sap_csv(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.SAP_XLSX:
            records_created, errors = parse_sap_xlsx(raw, batch, tenant)
        elif source == RawPayload.SourceSystem.SAP_ODATA:
            records_created, errors = parse_sap_odata(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.SAP_IDOC:
            records_created, errors = parse_sap_idoc(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.UTILITY_IN_CSV:
            records_created, errors = parse_utility_in(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.UTILITY_UK_CSV:
            records_created, errors = parse_utility_uk(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.TRAVEL_CSV:
            records_created, errors = parse_travel_csv(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.CONCUR_JSON:
            records_created, errors = parse_concur_json(raw, batch, content, tenant)
        elif source == RawPayload.SourceSystem.NAVAN_JSON:
            records_created, errors = parse_navan_json(raw, batch, content, tenant)

        # ── STATISTICAL OUTLIER DETECTION ACROSS BATCH (z-score > 3) ──
        records = bulk.records
        positive_vals = [float(r.normalized_value) for r in records if r.normalized_value and r.normalized_value > 0]
        
        if len(positive_vals) >= 5:
            import statistics
            mean_val = statistics.mean(positive_vals)
            stdev_val = statistics.stdev(positive_vals)
            if stdev_val > 0:
                for r in records:
                    if r.normalized_value and r.normalized_value > 0:
                        z = (float(r.normalized_value) - mean_val) / stdev_val
                        if z > 3:
                            r.is_suspicious = True
                            r.status = "FLAGGED"
                            errs = list(r.validation_errors or [])
                            msg = f"Statistical outlier: kgCO2e={float(r.normalized_value):.1f} is z={z:.1f} above batch mean {mean_val:.1f}"
                            if msg not in errs:
                                errs.append(msg)
                            r.validation_errors = errs
                            r.flag_reason = "; ".join(errs)

        # Update IngestionBatch details
        batch.row_count = len(records)
        batch.error_count = len(errors)
        batch.status = "SUCCESS" if len(errors) == 0 else "PARTIAL_SUCCESS"
        batch.save()

    return {"records_created": records_created, "errors_found": errors}


# --- SAP INGESTION PARSERS ---

def parse_sap_csv(raw, batch, content, tenant):
    lines = content.splitlines()
    reader = csv.reader(lines)
    headers = next(reader, None)
    if not headers:
        return 0, ["Empty CSV file"]
    
    resolved_headers = [ColumnNormalizer.resolve_header(h) for h in headers]
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    
    airports = Airport.objects.all()
    coords_map = {ap.iata_code.upper(): (float(ap.latitude), float(ap.longitude)) for ap in airports}
    for key, val in [("DEL", (28.5665, 77.1031)), ("BOM", (19.0896, 72.8656)), ("LHR", (51.4775, -0.4614)), ("SFO", (37.6213, -122.3790)), ("JFK", (40.6413, -73.7781)), ("SIN", (1.3644, 103.9915))]:
        if key not in coords_map:
            coords_map[key] = val
    country_map = {ap.iata_code.upper(): ap.country for ap in airports}
    # ──────────────────────────────────────────────
    
    for idx, row in enumerate(reader):
        if not row or not any(row):
            continue
        row_dict = dict(zip(resolved_headers, row))
        
        try:
            txn_id = row_dict.get('txn_id') or f"csv-{raw.id.hex[:6]}-{idx}"
            mvt = row_dict.get('movement_type', '101')
            qty = sanitize_decimal(row_dict.get('quantity', '0'))
            unit = row_dict.get('unit', 'L')
            mat = row_dict.get('material_number', '')
            plant = row_dict.get('plant', 'PL01')
            asset = row_dict.get('fixed_asset', '')
            mat_desc = row_dict.get('material_description', 'Diesel Fuel Oil')
            
            p_date_str = (
                row_dict.get('posting_date') or 
                row_dict.get('travel_date') or 
                row_dict.get('booking_date') or 
                row_dict.get('departure_date') or 
                row_dict.get('date') or 
                row_dict.get('transaction_date') or 
                row_dict.get('checkin_date') or 
                row_dict.get('created_date') or 
                ''
            )
            p_date = parse_date_flexible(p_date_str)
            
            is_suspicious = False
            validation_errors = []
            
            exp_type = str(row_dict.get('expense_type', '')).upper().strip()
            
            if exp_type in ('AIR', 'HOTEL', 'GROUND', 'PERDIEM'):
                # Process as travel
                if exp_type == 'AIR':
                    origin = str(row_dict.get('departure_airport', 'DEL')).upper().strip()
                    dest = str(row_dict.get('arrival_airport', 'BOM')).upper().strip()
                    cabin = str(row_dict.get('cabin_class', 'Economy')).strip()
                    
                    dist_km = None
                    if row_dict.get('distance_km'):
                        try:
                            dist_km = float(row_dict.get('distance_km'))
                        except ValueError:
                            pass
                    
                    res = normalize_flight(dist_km, origin, dest, cabin, coords_map)
                    if res.success:
                        norm_qty = Decimal(str(res.canonical_qty))
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal(str(res.kgco2e))
                        emission_factor = Decimal(str(res.emission_factor))
                        emission_factor_source = res.emission_factor_source
                        scope_str = str(res.scope)
                        category = "flights"
                        if res.flag:
                            validation_errors.append(res.flag)
                    else:
                        norm_qty = Decimal('0.0000')
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal('0.0000')
                        emission_factor = Decimal('0.0000')
                        emission_factor_source = ""
                        scope_str = "3"
                        category = "flights"
                        is_suspicious = True
                        validation_errors.append(res.flag)
                    
                    plant = origin
                    country = country_map.get(origin, 'IN')
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL
                    
                elif exp_type == 'HOTEL':
                    hotel_name = row_dict.get('hotel_name', 'Default Hotel')
                    hotel_country = row_dict.get('hotel_country', 'IN').upper().strip() or 'IN'
                    nights_str = row_dict.get('nights') or row_dict.get('hotel_nights')
                    try:
                        nights = int(nights_str) if nights_str else 1
                    except ValueError:
                        nights = 1
                    
                    qty = Decimal(str(nights))
                    unit = 'nights'
                    res = normalize_hotel(nights, hotel_country)
                    if res.success:
                        norm_qty = Decimal(str(res.canonical_qty))
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal(str(res.kgco2e))
                        emission_factor = Decimal(str(res.emission_factor))
                        emission_factor_source = res.emission_factor_source
                        scope_str = str(res.scope)
                        category = "hotel"
                    else:
                        norm_qty = Decimal('0.0000')
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal('0.0000')
                        emission_factor = Decimal('0.0000')
                        emission_factor_source = ""
                        scope_str = "3"
                        category = "hotel"
                        is_suspicious = True
                        validation_errors.append(res.flag)
                    
                    plant = hotel_name
                    country = hotel_country
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL
                    
                elif exp_type == 'GROUND':
                    ground_type = row_dict.get('ground_transport_type', 'Metro')
                    dist_str = row_dict.get('ground_distance_km') or row_dict.get('distance_km')
                    try:
                        distance_km = float(dist_str) if dist_str else 0.0
                    except ValueError:
                        distance_km = 0.0
                    
                    qty = Decimal(str(distance_km))
                    unit = 'km'
                    res = normalize_ground(distance_km, ground_type)
                    if res.success:
                        norm_qty = Decimal(str(res.canonical_qty))
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal(str(res.kgco2e))
                        emission_factor = Decimal(str(res.emission_factor))
                        emission_factor_source = res.emission_factor_source
                        scope_str = str(res.scope)
                        category = "ground"
                        if res.flag:
                            validation_errors.append(res.flag)
                    else:
                        norm_qty = Decimal('0.0000')
                        norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                        normalized_value = Decimal('0.0000')
                        emission_factor = Decimal('0.0000')
                        emission_factor_source = ""
                        scope_str = "3"
                        category = "ground"
                        is_suspicious = True
                        validation_errors.append(res.flag)
                    
                    plant = ground_type
                    country = 'IN'
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL
                    
                elif exp_type == 'PERDIEM':
                    qty = Decimal('0.0000')
                    unit = 'USD'
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = "Per Diem Exclusion"
                    scope_str = "3"
                    category = "per_diem"
                    is_suspicious = True
                    validation_errors.append("Per diem allowance — not a travel emission. Isolated from accounting.")
                    plant = "Per Diem"
                    country = "IN"
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL
            else:
                # Standard SAP Movement type and fuel calculation logic
                if mvt == '241':
                    scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
                    if unit.upper() in ('GAL', 'GALLON', 'GALLONS'):
                        norm_qty = qty * Decimal('3.78541')
                        norm_unit = EmissionRecord.NormalizedUnit.LITERS
                    else:
                        norm_qty = qty
                        norm_unit = EmissionRecord.NormalizedUnit.LITERS
                    if not asset:
                        is_suspicious = True
                        validation_errors.append("Scope 1 goods issue (241) lacks a fixed asset ID.")
                elif mvt == '101':
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                    norm_qty = qty
                    norm_unit = EmissionRecord.NormalizedUnit.METRIC_TONS if unit.upper() in ('TO', 'TON', 'TONS') else EmissionRecord.NormalizedUnit.LITERS
                elif mvt in ('313', '315'):
                    scope = EmissionRecord.ScopeCategory.EXCLUDED_LOGISTICS
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.LITERS
                else:
                    scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                    norm_qty = qty
                    norm_unit = EmissionRecord.NormalizedUnit.LITERS
                
                if qty < 0:
                    is_suspicious = True
                    validation_errors.append("Detected SAP reversal or credit memo transaction (negative quantity).")

                # Perform DEFRA 2023 conversions
                res = normalize_fuel(float(qty), unit, mat_desc)
                if res.success:
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "fuel"
                    if res.flag:
                        validation_errors.append(res.flag)
                else:
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "1"
                    category = "fuel"
                    is_suspicious = True
                    validation_errors.append(res.flag)

                plant_countries = {'PL01': 'IN', 'PL02': 'IN', 'PL04': 'GB', 'PL05': 'US'}
                country = plant_countries.get(plant.upper(), 'IN')

            dedup_key = f"sap_csv_{txn_id}_{mvt}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 1,
                unique_transaction_id=txn_id,
                scope_category=scope,
                start_date=p_date,
                end_date=p_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=p_date,
                period_end=p_date,
                source_row_id=hashlib_row_id(row_dict),
                raw_data=row_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=plant,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested from SAP Procurement CSV."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Row {idx+1} error: {str(e)}")
            
    return len(records), errors

def parse_sap_xlsx(raw, batch, tenant):
    try:
        if batch and batch.raw_file:
            # Open file-like object directly from active storage engine (S3 / Local / Memory)
            batch.raw_file.open('rb')
            wb = load_workbook(batch.raw_file, data_only=True)
        else:
            raise ValueError("No raw_file linked to batch.")
    except Exception as e:
        # Fallback to local default file path if remote load fails
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        xlsx_path = os.path.join(base_dir, "sample_data", "sap", "sap_fuel_consumption.xlsx")
        if not os.path.exists(xlsx_path):
            return 0, [f"Excel file does not exist locally ({xlsx_path}) and remote load failed: {str(e)}"]
        wb = load_workbook(xlsx_path, data_only=True)
        
    sheet = wb.active
    
    rows = list(sheet.rows)
    if not rows:
        return 0, ["Empty Excel sheet"]
        
    headers = [cell.value for cell in rows[0] if cell.value is not None and str(cell.value).strip() != ""]
    resolved_headers = [ColumnNormalizer.resolve_header(h) for h in headers]
    
    plant_lookup = {'PL01': 'IN', 'PL02': 'IN', 'PL04': 'GB', 'PL05': 'US'}
    if "Plant_Master" in wb.sheetnames:
        pm_sheet = wb["Plant_Master"]
        for p_row in pm_sheet.iter_rows(min_row=2, values_only=True):
            if p_row[0]:
                plant_lookup[str(p_row[0]).upper()] = str(p_row[2] or 'IN')
                
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    for idx, row in enumerate(rows[1:]):
        vals = [cell.value for cell in row[:len(headers)]]
        if not any(vals):
            continue
        row_dict = dict(zip(resolved_headers, vals))
        
        try:
            txn_id = str(row_dict.get('txn_id') or f"xlsx-{raw.id.hex[:6]}-{idx}")
            mvt = str(row_dict.get('movement_type', '241'))
            qty = sanitize_decimal(row_dict.get('quantity', '0'))
            unit = str(row_dict.get('unit', 'L'))
            plant = str(row_dict.get('plant', 'PL01')).upper()
            asset = str(row_dict.get('fixed_asset', ''))
            mat_desc = str(row_dict.get('material_description', 'Diesel Fuel Oil'))
            
            p_date_raw = (
                row_dict.get('posting_date') or 
                row_dict.get('travel_date') or 
                row_dict.get('booking_date') or 
                row_dict.get('departure_date') or 
                row_dict.get('date') or 
                row_dict.get('transaction_date') or 
                row_dict.get('checkin_date') or 
                row_dict.get('created_date')
            )
            if isinstance(p_date_raw, datetime):
                p_date = timezone.make_aware(p_date_raw, pytz.utc)
            elif p_date_raw:
                p_date = parse_date_flexible(str(p_date_raw))
            else:
                p_date = timezone.now()
                
            is_suspicious = False
            validation_errors = []
            
            if qty < 0:
                is_suspicious = True
                validation_errors.append("Detected SAP reversal or credit memo transaction (negative quantity).")
                
            if unit.upper() in ('KG', 'KILOGRAM', 'KILOGRAMS', 'TO'):
                is_suspicious = True
                validation_errors.append(f"Standardized baseline expects volume (Liters), but mass unit '{unit}' was provided.")
                
            if unit.upper() in ('GAL', 'GALLON', 'GALLONS'):
                norm_qty = qty * Decimal('3.78541')
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
            else:
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.LITERS

            res = normalize_fuel(float(qty), unit, mat_desc)
            if res.success:
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "fuel"
                if res.flag:
                    validation_errors.append(res.flag)
            else:
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = ""
                scope_str = "1"
                category = "fuel"
                is_suspicious = True
                validation_errors.append(res.flag)

            scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
            country = plant_lookup.get(plant, 'IN')
            dedup_key = f"sap_xlsx_{txn_id}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 2,
                unique_transaction_id=txn_id,
                scope_category=scope,
                start_date=p_date,
                end_date=p_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=p_date,
                period_end=p_date,
                source_row_id=hashlib_row_id(row_dict),
                raw_data=row_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=plant,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested from SAP Fuel XLSX."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Row {idx+2} error: {str(e)}")
            
    return len(records), errors

def parse_sap_odata(raw, batch, content, tenant):
    payload = json.loads(content)
    results = payload.get('d', {}).get('results', [])
    
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    for idx, item in enumerate(results):
        try:
            txn_id = item.get('MaterialDocument')
            mvt = item.get('GoodsMovementType')
            qty_str = item.get('QuantityInEntryUnit')
            qty = sanitize_decimal(qty_str)
            unit = item.get('EntryUnit')
            plant = item.get('Plant', 'PL01').upper()
            asset = item.get('MasterFixedAsset', '')
            posting_date_str = item.get('PostingDate')
            mat_desc = item.get('MaterialDescription', 'Diesel Fuel Oil')
            
            p_date = parse_microsoft_epoch(posting_date_str)
            
            is_suspicious = False
            validation_errors = []
            
            if mvt == '241':
                scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
                if not asset:
                    is_suspicious = True
                    validation_errors.append("Scope 1 goods issue (241) lacks a fixed asset ID.")
            elif mvt == '101':
                scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.METRIC_TONS if unit.upper() in ('TO', 'TON', 'TONS') else EmissionRecord.NormalizedUnit.LITERS
            elif mvt in ('313', '315'):
                scope = EmissionRecord.ScopeCategory.EXCLUDED_LOGISTICS
                norm_qty = Decimal('0.0000')
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
            else:
                scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
                
            if qty < 0:
                is_suspicious = True
                validation_errors.append("Detected SAP reversal or credit memo transaction (negative quantity).")

            res = normalize_fuel(float(qty), unit, mat_desc)
            if res.success:
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "fuel"
                if res.flag:
                    validation_errors.append(res.flag)
            else:
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = ""
                scope_str = "1"
                category = "fuel"
                is_suspicious = True
                validation_errors.append(res.flag)

            plant_countries = {'PL01': 'IN', 'PL02': 'IN', 'PL04': 'GB', 'PL05': 'US'}
            country = plant_countries.get(plant, 'IN')
            dedup_key = f"sap_odata_{txn_id}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx,
                unique_transaction_id=txn_id,
                scope_category=scope,
                start_date=p_date,
                end_date=p_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=p_date,
                period_end=p_date,
                source_row_id=hashlib_row_id(item),
                raw_data=item,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=plant,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested from SAP OData API payload."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Result element {idx} error: {str(e)}")
            
    return len(records), errors

def parse_sap_idoc(raw, batch, content, tenant):
    content_stripped = content.strip()
    if content_stripped.startswith('<'):
        return parse_sap_idoc_xml(raw, batch, content, tenant)
    else:
        return parse_sap_idoc_flat(raw, batch, content, tenant)

def parse_sap_idoc_xml(raw, batch, content, tenant):
    class SafeParser(ET.XMLParser):
        def doctype(self, name, pubid, system):
            raise ValidationError("XML Document Type Definition (DTD) is forbidden for security reasons.")
            
    root = ET.fromstring(content, parser=SafeParser())
    
    testrun_elem = root.find('.//TESTRUN')
    is_testrun = testrun_elem is not None and testrun_elem.text == 'X'
    
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    items = root.findall('.//E1BP2017_GM_ITEM_CREATE')
    doc_num_elem = root.find('.//DOCNUM')
    doc_num = doc_num_elem.text if doc_num_elem is not None else f"idoc-{raw.id.hex[:6]}"
    
    pstng_date_elem = root.find('.//PSTNG_DATE')
    if pstng_date_elem is None or not pstng_date_elem.text:
        pstng_date_elem = root.find('.//BLDAT')
    if pstng_date_elem is None or not pstng_date_elem.text:
        pstng_date_elem = root.find('.//BUDAT')
    if pstng_date_elem is None or not pstng_date_elem.text:
        pstng_date_elem = root.find('.//CREATION_DATE')
        
    if pstng_date_elem is not None and pstng_date_elem.text:
        d_str = pstng_date_elem.text
        p_date = parse_date_flexible(d_str)
    else:
        p_date = timezone.now()

    for idx, item in enumerate(items):
        try:
            mat = item.find('MATERIAL').text if item.find('MATERIAL') is not None else ''
            plant = item.find('PLANT').text if item.find('PLANT') is not None else 'PL01'
            mvt = item.find('MOVE_TYPE').text if item.find('MOVE_TYPE') is not None else '101'
            qty_str = item.find('ENTRY_QNT').text if item.find('ENTRY_QNT') is not None else '0'
            qty = sanitize_decimal(qty_str)
            unit = item.find('ENTRY_UOM').text if item.find('ENTRY_UOM') is not None else 'L'
            mat_desc = "Diesel Fuel Oil"
            
            is_suspicious = is_testrun
            validation_errors = []
            if is_testrun:
                validation_errors.append("SAP TESTRUN flag is set. This simulated transaction was not committed in SAP.")

            if mvt == '241':
                scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
            elif mvt == '101':
                scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.METRIC_TONS if unit.upper() in ('TO', 'TON', 'TONS') else EmissionRecord.NormalizedUnit.LITERS
            else:
                scope = EmissionRecord.ScopeCategory.EXCLUDED_LOGISTICS
                norm_qty = Decimal('0.0000')
                norm_unit = EmissionRecord.NormalizedUnit.LITERS

            res = normalize_fuel(float(qty), unit, mat_desc)
            if res.success:
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "fuel"
                if res.flag:
                    validation_errors.append(res.flag)
            else:
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = ""
                scope_str = "1"
                category = "fuel"
                is_suspicious = True
                validation_errors.append(res.flag)
                
            plant_countries = {'PL01': 'IN', 'PL02': 'IN', 'PL04': 'GB', 'PL05': 'US'}
            country = plant_countries.get(plant.upper(), 'IN')
            dedup_key = f"sap_idoc_{doc_num}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            # Reconstruct XML row representation
            item_xml_dict = {elem.tag: elem.text for elem in item}

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 1,
                unique_transaction_id=doc_num,
                scope_category=scope,
                start_date=p_date,
                end_date=p_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=p_date,
                period_end=p_date,
                source_row_id=hashlib_row_id(item_xml_dict),
                raw_data=item_xml_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=plant,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested from legacy SAP XML IDoc payload."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"XML Segment index {idx} error: {str(e)}")
            
    return len(records), errors

def parse_sap_idoc_flat(raw, batch, content, tenant):
    lines = content.splitlines()
    
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    items_list = []
    current_item = None
    
    is_testrun = False
    doc_num = f"idoc-{raw.id.hex[:6]}"
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        parts = line.split()
        if not parts:
            continue
        seg_name = parts[0]
        
        if seg_name == 'E2EDK01005':
            if 'NB' in parts:
                nb_idx = parts.index('NB')
                if len(parts) > nb_idx + 1:
                    doc_num = parts[nb_idx + 1]
        elif seg_name == 'E2EDP01005':
            if current_item:
                items_list.append(current_item)
            current_item = {
                'qty': sanitize_decimal(parts[3]) if len(parts) > 3 else Decimal('0'),
                'unit': parts[4] if len(parts) > 4 else 'L',
                'date': timezone.now(),
                'desc': 'Generator Diesel Drums 200L',
                'plant': 'PL01',
                'mvt': '101'
            }
        elif seg_name == 'E2EDP20' and current_item:
            if len(parts) > 2:
                current_item['date'] = parse_date_flexible(parts[2])
        elif seg_name == 'E2EDP19001' and current_item:
            if len(parts) > 2:
                desc_text = ' '.join(parts[2:])
                if desc_text and not desc_text.isdigit() and len(desc_text) > 3:
                    current_item['desc'] = desc_text

    if current_item:
        items_list.append(current_item)

    for idx, item in enumerate(items_list):
        try:
            qty = item['qty']
            unit = item['unit']
            p_date = item['date']
            mat_desc = item['desc']
            plant = item['plant']
            mvt = item['mvt']
            
            is_suspicious = is_testrun
            validation_errors = []
            
            if mvt == '241':
                scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.LITERS
            elif mvt == '101':
                scope = EmissionRecord.ScopeCategory.SCOPE_3_PROCUREMENT
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.METRIC_TONS if unit.upper() in ('TO', 'TON', 'TONS', 'KG', 'KGM') else EmissionRecord.NormalizedUnit.LITERS
            else:
                scope = EmissionRecord.ScopeCategory.EXCLUDED_LOGISTICS
                norm_qty = Decimal('0.0000')
                norm_unit = EmissionRecord.NormalizedUnit.LITERS

            res = normalize_fuel(float(qty), unit, mat_desc)
            if res.success:
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "fuel"
                if res.flag:
                    validation_errors.append(res.flag)
            else:
                # Fallback for HFC R-410A refrigerant factor
                if 'refrigerant' in mat_desc.lower() or 'r-410a' in mat_desc.lower():
                    emission_factor = Decimal('2088.0')
                    normalized_value = qty * emission_factor
                    emission_factor_source = "DEFRA 2023 HFC Refrigerant Factors"
                    scope_str = "1"
                    category = "refrigerant"
                    scope = EmissionRecord.ScopeCategory.SCOPE_1_STATIONARY
                else:
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "1"
                    category = "fuel"
                    is_suspicious = True
                    validation_errors.append(res.flag)
                
            plant_countries = {'PL01': 'IN', 'PL02': 'IN', 'PL04': 'GB', 'PL05': 'US'}
            country = plant_countries.get(plant.upper(), 'IN')
            dedup_key = f"sap_idoc_flat_{doc_num}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            item_dict = {
                'material': mat_desc,
                'quantity': str(qty),
                'unit': unit,
                'posting_date': p_date.strftime('%Y-%m-%d'),
                'plant': plant,
                'doc_num': doc_num
            }

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 1,
                unique_transaction_id=doc_num,
                scope_category=scope,
                start_date=p_date,
                end_date=p_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=p_date,
                period_end=p_date,
                source_row_id=hashlib_row_id(item_dict),
                raw_data=item_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=plant,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested from legacy flat SAP IDoc segment."
            )
            records.append(activity)

        except Exception as e:
            errors.append(f"Flat IDoc element {idx} error: {str(e)}")

    return len(records), errors


# --- UTILITY INGESTION PARSERS ---

def parse_utility_in(raw, batch, content, tenant):
    lines = content.splitlines()
    reader = csv.reader(lines)
    headers = next(reader, None)
    if not headers:
        return 0, ["Empty Indian Utility CSV"]
        
    resolved_headers = [ColumnNormalizer.resolve_header(h) for h in headers]
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    for idx, row in enumerate(reader):
        if not row or not any(row):
            continue
        row_dict = dict(zip(resolved_headers, row))
        try:
            meter_id = row_dict.get('meter_id') or row_dict.get('txn_id') or 'IN-METER-01'
            qty = sanitize_decimal(row_dict.get('quantity') or row_dict.get('consumption') or row_dict.get('qty') or '0')
            unit = row_dict.get('unit', 'kWh')
            
            start_str = row_dict.get('billing_start') or row_dict.get('posting_date') or row_dict.get('start_date') or row_dict.get('period_start')
            end_str = row_dict.get('billing_end') or row_dict.get('end_date') or row_dict.get('period_end')
            
            s_date = parse_date_flexible(start_str)
            e_date = parse_date_flexible(end_str)
            
            is_suspicious = False
            validation_errors = []
            
            if e_date <= s_date:
                is_suspicious = True
                validation_errors.append(f"Invalid billing duration: billing end date ({end_str}) is before start date ({start_str}).")
            
            # Solar export overrides
            if qty < 0:
                scope = EmissionRecord.ScopeCategory.EXCLUDED_AVOIDED
                norm_qty = abs(qty)
                norm_unit = EmissionRecord.NormalizedUnit.KWH
                validation_errors.append("Negative usage value representing net-metered solar export to the grid. Isolated to Avoided Emissions.")
                is_suspicious = True
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = "Avoided Emissions Offset"
                scope_str = "3"
                category = "exported_renewable_energy"
            else:
                scope = EmissionRecord.ScopeCategory.SCOPE_2_ELECTRICITY
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.KWH
                
                # Perform DEFRA Location based calculation
                res = normalize_electricity(float(qty), unit, "IN")
                if res.success:
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "electricity"
                else:
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "2"
                    category = "electricity"
                    is_suspicious = True
                    validation_errors.append(res.flag)
                
            delta = e_date - s_date
            total_days = max(delta.days, 1)
            daily_rate = norm_qty / Decimal(str(total_days))
            daily_emissions = normalized_value / Decimal(str(total_days))
            
            current_day = s_date
            months_traversed = {}
            for day_idx in range(total_days):
                day_key = current_day.strftime("%Y-%m")
                months_traversed.setdefault(day_key, []).append(current_day)
                current_day += timedelta(days=1)
                
            for month_key, days_list in months_traversed.items():
                m_start = days_list[0]
                m_end = days_list[-1].replace(hour=23, minute=59, second=59)
                m_days = len(days_list)
                m_qty = daily_rate * Decimal(str(m_days))
                m_emissions = daily_emissions * Decimal(str(m_days))
                
                dedup_key = f"utility_in_{meter_id}_{month_key}_{idx}"
                status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"
                
                activity = EmissionRecord.objects.create(
                    tenant=tenant,
                    raw_payload=raw,
                    batch=batch,
                    source_row_index=idx + 1,
                    unique_transaction_id=meter_id,
                    scope_category=scope,
                    start_date=m_start,
                    end_date=m_end,
                    raw_quantity=qty * Decimal(str(m_days)) / Decimal(str(total_days)),
                    raw_unit=unit,
                    normalized_quantity=m_qty,
                    normalized_unit=norm_unit,
                    scope=scope_str,
                    category=category,
                    activity_value=qty * Decimal(str(m_days)) / Decimal(str(total_days)),
                    activity_unit=unit,
                    normalized_value=m_emissions,
                    normalized_value_unit="kgCO2e",
                    emission_factor=emission_factor,
                    emission_factor_source=emission_factor_source,
                    period_start=m_start,
                    period_end=m_end,
                    source_row_id=hashlib_row_id(row_dict),
                    raw_data=row_dict,
                    status=status,
                    flag_reason="; ".join(validation_errors),
                    resolved_facility_id=f"Facility-IN-{meter_id}",
                    resolved_facility_country='IN',
                    is_suspicious=is_suspicious,
                    validation_errors=validation_errors,
                    deduplication_key=dedup_key
                )

                AuditLog.objects.create(
                    activity=activity,
                    action=AuditLog.AuditAction.CREATE,
                    changed_by=raw.ingested_by,
                    performed_by=user_performer,
                    reason=f"Auto-ingested & linearly prorated utility billing block segment for {month_key} ({m_days} days)."
                )
                records.append(activity)
                
        except Exception as e:
            errors.append(f"Row {idx+1} error: {str(e)}")
            
    return len(records), errors

def parse_utility_uk(raw, batch, content, tenant):
    lines = content.splitlines()
    reader = csv.reader(lines)
    headers = next(reader, None)
    if not headers:
        return 0, ["Empty UK smart meter CSV"]
        
    resolved_headers = [ColumnNormalizer.resolve_header(h) for h in headers]
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    # ──────────────────────────────────────────────
    
    fac_tz = {'METER-UK-01': 'Europe/London', 'METER-UK-02': 'Europe/London'}
    
    for idx, row in enumerate(reader):
        if not row or not any(row):
            continue
        row_dict = dict(zip(resolved_headers, row))
        
        try:
            meter_id = row_dict.get('meter_id') or row_dict.get('txn_id') or 'METER-UK-01'
            qty = sanitize_decimal(row_dict.get('quantity') or row_dict.get('consumption_kwh') or row_dict.get('qty') or '0')
            unit = row_dict.get('unit', 'kWh')
            
            start_str = row_dict.get('billing_start') or row_dict.get('posting_date') or row_dict.get('start_date') or row_dict.get('period_start')
            end_str = row_dict.get('billing_end') or row_dict.get('end_date') or row_dict.get('period_end') or start_str
            
            tz_name = fac_tz.get(meter_id, 'Europe/London')
            local_tz = pytz.timezone(tz_name)
            
            s_dt = parse_date_flexible(start_str)
            e_dt = parse_date_flexible(end_str)
            
            naive_start = timezone.make_naive(s_dt, pytz.utc) if timezone.is_aware(s_dt) else s_dt
            naive_end = timezone.make_naive(e_dt, pytz.utc) if timezone.is_aware(e_dt) else e_dt
            
            local_start = local_tz.localize(naive_start, is_dst=True)
            utc_start = local_start.astimezone(pytz.utc)
            
            local_end = local_tz.localize(naive_end, is_dst=True)
            utc_end = local_end.astimezone(pytz.utc)
            
            is_suspicious = False
            validation_errors = []
            if unit.upper() in ('MWH', 'MEGAWATT-HOURS', 'MEGAWATT-HOUR'):
                norm_qty = qty * Decimal('1000.0000')
                norm_unit = EmissionRecord.NormalizedUnit.KWH
            else:
                norm_qty = qty
                norm_unit = EmissionRecord.NormalizedUnit.KWH
                
            reading_type = str(row_dict.get('reading_type') or row_dict.get('is_estimated') or '').lower()
            if reading_type == 'true' or not meter_id:
                is_suspicious = True
                validation_errors.append("Calculated based on estimated utility readings; potential grid margin deviations.")

            res = normalize_electricity(float(qty), unit, "GB")
            if res.success:
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "electricity"
            else:
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = ""
                scope_str = "2"
                category = "electricity"
                is_suspicious = True
                validation_errors.append(res.flag)
                
            dedup_key = f"utility_uk_{meter_id}_{utc_start.strftime('%Y%m%d%H%M')}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 1,
                unique_transaction_id=meter_id,
                scope_category=EmissionRecord.ScopeCategory.SCOPE_2_ELECTRICITY,
                start_date=utc_start,
                end_date=utc_end,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=utc_start,
                period_end=utc_end,
                source_row_id=hashlib_row_id(row_dict),
                raw_data=row_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=f"Facility-UK-{meter_id}",
                resolved_facility_country='GB',
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested UK smart meter time-series interval feed."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Row {idx+1} error: {str(e)}")
            
    return len(records), errors


# --- TRAVEL INGESTION PARSERS ---

def parse_travel_csv(raw, batch, content, tenant):
    lines = content.splitlines()
    reader = csv.reader(lines)
    headers = next(reader, None)
    if not headers:
        return 0, ["Empty Corporate Travel CSV"]
        
    resolved_headers = [ColumnNormalizer.resolve_header(h) for h in headers]
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    
    airports = Airport.objects.all()
    coords_map = {ap.iata_code.upper(): (float(ap.latitude), float(ap.longitude)) for ap in airports}
    for key, val in [("DEL", (28.5665, 77.1031)), ("BOM", (19.0896, 72.8656)), ("LHR", (51.4775, -0.4614)), ("SFO", (37.6213, -122.3790)), ("JFK", (40.6413, -73.7781)), ("SIN", (1.3644, 103.9915))]:
        if key not in coords_map:
            coords_map[key] = val
            
    country_map = {ap.iata_code.upper(): ap.country for ap in airports}
    # ──────────────────────────────────────────────
    
    for idx, row in enumerate(reader):
        if not row or not any(row):
            continue
        row_dict = dict(zip(resolved_headers, row))
        
        try:
            trip_id = row_dict.get('txn_id') or row_dict.get('trip_id') or f"travel-{raw.id.hex[:6]}-{idx}"
            exp_type = str(row_dict.get('expense_type', 'AIR')).upper().strip()
            
            date_str = (
                row_dict.get('posting_date') or 
                row_dict.get('travel_date') or 
                row_dict.get('booking_date') or 
                row_dict.get('departure_date') or 
                row_dict.get('date') or 
                row_dict.get('transaction_date') or 
                row_dict.get('checkin_date') or 
                row_dict.get('created_date') or 
                ''
            )
            t_date = parse_date_flexible(date_str)
            
            is_suspicious = False
            validation_errors = []
            
            if exp_type == 'AIR':
                origin = str(row_dict.get('departure_airport', 'DEL')).upper().strip()
                dest = str(row_dict.get('arrival_airport', 'BOM')).upper().strip()
                cabin = str(row_dict.get('cabin_class', 'Economy')).strip()
                
                qty = sanitize_decimal(row_dict.get('quantity', '0'))
                if qty == 0 and row_dict.get('distance_km'):
                    qty = sanitize_decimal(row_dict.get('distance_km'))
                unit = row_dict.get('unit') or 'passenger-km'

                dist_km = None
                if row_dict.get('distance_km'):
                    try:
                        dist_km = float(row_dict.get('distance_km'))
                    except ValueError:
                        pass

                res = normalize_flight(dist_km, origin, dest, cabin, coords_map)
                if res.success:
                    norm_qty = Decimal(str(res.canonical_qty))
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "flights"
                    if res.flag:
                        validation_errors.append(res.flag)
                else:
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "3"
                    category = "flights"
                    is_suspicious = True
                    validation_errors.append(res.flag)

                if cabin.lower() in ('business', 'first', 'j', 'f') and norm_qty < Decimal('300'):
                    is_suspicious = True
                    validation_errors.append("Suspicious Travel: Business or First Class selected for a short flight (< 300 km).")
                
                end_date = t_date + timedelta(hours=2)
                facility_id = origin
                country = country_map.get(origin, 'IN')
                audit_note = "Auto-ingested flight record with dynamic Haversine coordinates + 8% detour uplift."

            elif exp_type == 'HOTEL':
                hotel_name = row_dict.get('hotel_name', 'Default Hotel')
                hotel_city = row_dict.get('hotel_city', 'Unknown City')
                hotel_country = row_dict.get('hotel_country', 'IN').upper().strip() or 'IN'
                
                nights_str = row_dict.get('nights') or row_dict.get('hotel_nights')
                try:
                    nights = int(nights_str) if nights_str else 1
                except ValueError:
                    nights = 1
                    
                qty = Decimal(str(nights))
                unit = 'nights'
                
                checkout_str = row_dict.get('checkout_date')
                if checkout_str:
                    end_date = parse_date_flexible(checkout_str)
                else:
                    end_date = t_date + timedelta(days=nights)
                    
                res = normalize_hotel(nights, hotel_country)
                if res.success:
                    norm_qty = Decimal(str(res.canonical_qty))
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "hotel"
                else:
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "3"
                    category = "hotel"
                    is_suspicious = True
                    validation_errors.append(res.flag)
                    
                facility_id = hotel_name
                country = hotel_country
                audit_note = "Auto-ingested hotel lodging stay record with country-specific DEFRA emission factors."

            elif exp_type == 'GROUND':
                ground_type = row_dict.get('ground_transport_type', 'Metro')
                dist_str = row_dict.get('ground_distance_km') or row_dict.get('distance_km')
                try:
                    distance_km = float(dist_str) if dist_str else 0.0
                except ValueError:
                    distance_km = 0.0
                    
                qty = Decimal(str(distance_km))
                unit = 'km'
                end_date = t_date + timedelta(hours=1)
                
                res = normalize_ground(distance_km, ground_type)
                if res.success:
                    norm_qty = Decimal(str(res.canonical_qty))
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "ground"
                    if res.flag:
                        validation_errors.append(res.flag)
                else:
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "3"
                    category = "ground"
                    is_suspicious = True
                    validation_errors.append(res.flag)
                    
                facility_id = ground_type
                country = 'IN'
                audit_note = "Auto-ingested ground transportation record with type-specific DEFRA emission factors."

            elif exp_type == 'PERDIEM':
                qty = Decimal('0.0000')
                unit = 'USD'
                end_date = t_date
                
                is_suspicious = True
                validation_errors.append("Per diem allowance — not a travel emission. Isolated from accounting.")
                
                norm_qty = Decimal('0.0000')
                norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = "Per Diem Exclusion"
                scope_str = "3"
                category = "per_diem"
                
                facility_id = "Per Diem"
                country = "IN"
                audit_note = "Auto-ingested per diem allowance; flagged and excluded from accounting."
                
            else:
                raise ValueError(f"Unknown expense_type '{exp_type}'")

            dedup_key = f"travel_csv_{trip_id}_{exp_type.lower()}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"

            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx + 1,
                unique_transaction_id=trip_id,
                scope_category=EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL,
                start_date=t_date,
                end_date=end_date,
                raw_quantity=qty,
                raw_unit=unit,
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=qty,
                activity_unit=unit,
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=t_date,
                period_end=end_date,
                source_row_id=hashlib_row_id(row_dict),
                raw_data=row_dict,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=facility_id,
                resolved_facility_country=country,
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )

            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason=audit_note
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Row {idx+1} error: {str(e)}")
            
    return len(records), errors

def parse_concur_json(raw, batch, content, tenant):
    payload = json.loads(content)
    bookings = payload.get('ItineraryList', {}).get('Bookings', [])
    
    records = []
    errors = []
    
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    airports = Airport.objects.all()
    coords_map = {ap.iata_code.upper(): (float(ap.latitude), float(ap.longitude)) for ap in airports}
    for key, val in [("DEL", (28.5665, 77.1031)), ("BOM", (19.0896, 72.8656)), ("LHR", (51.4775, -0.4614)), ("SFO", (37.6213, -122.3790)), ("JFK", (40.6413, -73.7781)), ("SIN", (1.3644, 103.9915))]:
        if key not in coords_map:
            coords_map[key] = val
    country_map = {ap.iata_code.upper(): ap.country for ap in airports}
    
    for idx, booking in enumerate(bookings):
        try:
            trip_id = booking.get('Id')
            traveller = booking.get('Traveller', {})
            emp_id = traveller.get('EmployeeId')
            email = traveller.get('EmailAddress')
            
            segments = booking.get('Segments', [])
            
            for seg_idx, seg in enumerate(segments):
                seg_type = seg.get('SegmentType')
                if seg_type != 'AIR':
                    continue
                    
                dep = seg.get('Departure', {})
                arr = seg.get('Arrival', {})
                
                origin = dep.get('AirportCode')
                dest = arr.get('AirportCode')
                cabin = seg.get('CabinClass', 'Economy')
                
                dep_time_str = (
                    dep.get('DateTime') or 
                    booking.get('createdDate') or 
                    booking.get('created_date') or 
                    booking.get('PostingDate') or 
                    booking.get('Date')
                )
                t_date = parse_microsoft_epoch(dep_time_str)
                
                is_suspicious = False
                validation_errors = []
                
                if not emp_id:
                    is_suspicious = True
                    validation_errors.append("Audit Alert: Travel booked by non-employee (guest). Manual boundary check required.")

                res = normalize_flight(None, origin, dest, cabin, coords_map)
                if res.success:
                    norm_qty = Decimal(str(res.canonical_qty))
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                    scope_str = str(res.scope)
                    category = "flights"
                    if res.flag:
                        validation_errors.append(res.flag)
                else:
                    norm_qty = Decimal('0.0000')
                    norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                    scope_str = "3"
                    category = "flights"
                    is_suspicious = True
                    validation_errors.append(res.flag)

                dedup_key = f"travel_concur_{trip_id}_{seg_idx}"
                status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"
                
                activity = EmissionRecord.objects.create(
                    tenant=tenant,
                    raw_payload=raw,
                    batch=batch,
                    source_row_index=idx,
                    unique_transaction_id=trip_id,
                    scope_category=EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL,
                    start_date=t_date,
                    end_date=t_date + timedelta(hours=3),
                    raw_quantity=Decimal('0.0000'),
                    raw_unit='passenger-km',
                    normalized_quantity=norm_qty,
                    normalized_unit=norm_unit,
                    scope=scope_str,
                    category=category,
                    activity_value=Decimal('0.0000'),
                    activity_unit="passenger-km",
                    normalized_value=normalized_value,
                    normalized_value_unit="kgCO2e",
                    emission_factor=emission_factor,
                    emission_factor_source=emission_factor_source,
                    period_start=t_date,
                    period_end=t_date + timedelta(hours=3),
                    source_row_id=hashlib_row_id(seg),
                    raw_data=seg,
                    status=status,
                    flag_reason="; ".join(validation_errors),
                    resolved_facility_id=origin,
                    resolved_facility_country=country_map.get(origin.upper(), 'US'),
                    is_suspicious=is_suspicious,
                    validation_errors=validation_errors,
                    deduplication_key=dedup_key
                )
                
                AuditLog.objects.create(
                    activity=activity,
                    action=AuditLog.AuditAction.CREATE,
                    changed_by=raw.ingested_by,
                    performed_by=user_performer,
                    reason="Auto-ingested flight segment from nested Concur Itinerary payload."
                )
                records.append(activity)
                
        except Exception as e:
            errors.append(f"Booking element {idx} error: {str(e)}")
            
    return len(records), errors

def parse_navan_json(raw, batch, content, tenant):
    payload = json.loads(content)
    data = payload.get('navanTmcResponse', {}).get('data', [])
    
    records = []
    errors = []
    
    # ─── OPTIMIZED DB PRE-LOOKUPS OUTSIDE LOOP ───
    user_performer = User.objects.filter(email=raw.ingested_by).first()
    
    airports = Airport.objects.all()
    coords_map = {ap.iata_code.upper(): (float(ap.latitude), float(ap.longitude)) for ap in airports}
    for key, val in [("DEL", (28.5665, 77.1031)), ("BOM", (19.0896, 72.8656)), ("LHR", (51.4775, -0.4614)), ("SFO", (37.6213, -122.3790)), ("JFK", (40.6413, -73.7781)), ("SIN", (1.3644, 103.9915))]:
        if key not in coords_map:
            coords_map[key] = val
            
    country_map = {ap.iata_code.upper(): ap.country for ap in airports}
    # ──────────────────────────────────────────────
    
    for idx, booking in enumerate(data):
        try:
            trip_id = booking.get('bookingId')
            status_val = booking.get('bookingStatus', 'TICKETED').upper()
            passengers = booking.get('passengerInfo', [])
            
            is_guest = False
            emp_id = None
            if passengers:
                p_type = passengers[0].get('passengerType', 'EMPLOYEE')
                is_guest = (p_type == 'GUEST')
                emp_id = passengers[0].get('employeeId')
                
            flight = booking.get('flightDetails', {})
            origin = flight.get('originAirport')
            dest = flight.get('destinationAirport')
            cabin = flight.get('cabinClassCode', 'Y')
            
            dep_time_str = flight.get('departureDateTime') or booking.get('createdDate')
            t_date = parse_microsoft_epoch(dep_time_str)
            
            is_suspicious = False
            validation_errors = []
            
            if is_guest or not emp_id:
                is_suspicious = True
                validation_errors.append("Audit Alert: Travel booked by non-employee (guest). Manual boundary check required.")
                
            # Mutable webhook logic (Cancellations)
            if status_val == 'CANCELED':
                locked_prev = None
                prev_recs = EmissionRecord.objects.filter(unique_transaction_id=trip_id)
                for pr in prev_recs:
                    if pr.is_locked or pr.workflow_status == EmissionRecord.WorkflowStatus.LOCKED_FOR_AUDIT:
                        locked_prev = pr
                        break
                
                res = normalize_flight(None, origin, dest, cabin, coords_map)
                if res.success:
                    norm_qty = Decimal(str(res.canonical_qty))
                    normalized_value = Decimal(str(res.kgco2e))
                    emission_factor = Decimal(str(res.emission_factor))
                    emission_factor_source = res.emission_factor_source
                else:
                    norm_qty = Decimal('0.0000')
                    normalized_value = Decimal('0.0000')
                    emission_factor = Decimal('0.0000')
                    emission_factor_source = ""
                
                if locked_prev:
                    comp_date = timezone.now()
                    dedup_key = f"travel_navan_cancel_compensating_{trip_id}_{idx}"
                    status = "PENDING_REVIEW"
                    
                    activity = EmissionRecord.objects.create(
                        tenant=tenant,
                        raw_payload=raw,
                        batch=batch,
                        source_row_index=idx,
                        unique_transaction_id=trip_id,
                        scope_category=EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL,
                        start_date=comp_date,
                        end_date=comp_date + timedelta(hours=1),
                        raw_quantity=Decimal('0.0000'),
                        raw_unit='passenger-km',
                        normalized_quantity=-norm_qty,
                        normalized_unit=EmissionRecord.NormalizedUnit.PASSENGER_KM,
                        scope="3",
                        category="flights",
                        activity_value=Decimal('0.0000'),
                        activity_unit="passenger-km",
                        normalized_value=-normalized_value,
                        normalized_value_unit="kgCO2e",
                        emission_factor=emission_factor,
                        emission_factor_source=emission_factor_source,
                        period_start=comp_date,
                        period_end=comp_date + timedelta(hours=1),
                        source_row_id=hashlib_row_id(booking),
                        raw_data=booking,
                        status=status,
                        flag_reason="Navan Webhook: Cancellation received for previously LOCKED/audited transaction. Compensation offset written.",
                        resolved_facility_id=origin,
                        resolved_facility_country=country_map.get(origin.upper(), 'US') if origin else 'US',
                        is_suspicious=True,
                        validation_errors=["Navan Webhook: Cancellation received for previously LOCKED/audited transaction. Compensation offset written."],
                        deduplication_key=dedup_key
                    )
                    
                    AuditLog.objects.create(
                        activity=activity,
                        action=AuditLog.AuditAction.CREATE,
                        changed_by=raw.ingested_by,
                        performed_by=user_performer,
                        reason=f"Compensating ledger entry created to balance audited cancellation of trip {trip_id}."
                    )
                    records.append(activity)
                    continue
                else:
                    prev_records = EmissionRecord.objects.filter(unique_transaction_id=trip_id)
                    if prev_records.exists():
                        for r in prev_records:
                            r.workflow_status = EmissionRecord.WorkflowStatus.REJECTED
                            r.status = "REJECTED"
                            r.save()
                            
                            AuditLog.objects.create(
                                activity=r,
                                action=AuditLog.AuditAction.REJECT,
                                changed_by=raw.ingested_by,
                                performed_by=user_performer,
                                reason="Navan Webhook: Trip canceled prior to audit lock. Transaction rejected."
                            )
                    else:
                        # Create as a new rejected record!
                        res = normalize_flight(None, origin, dest, cabin, coords_map)
                        if res.success:
                            norm_qty = Decimal(str(res.canonical_qty))
                            normalized_value = Decimal(str(res.kgco2e))
                            emission_factor = Decimal(str(res.emission_factor))
                            emission_factor_source = res.emission_factor_source
                        else:
                            norm_qty = Decimal('0.0000')
                            normalized_value = Decimal('0.0000')
                            emission_factor = Decimal('0.0000')
                            emission_factor_source = ""
                        
                        dedup_key = f"travel_navan_{trip_id}_canceled_{idx}"
                        activity = EmissionRecord.objects.create(
                            tenant=tenant,
                            raw_payload=raw,
                            batch=batch,
                            source_row_index=idx,
                            unique_transaction_id=trip_id,
                            scope_category=EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL,
                            start_date=t_date,
                            end_date=t_date + timedelta(hours=3),
                            raw_quantity=Decimal('0.0000'),
                            raw_unit='passenger-km',
                            normalized_quantity=norm_qty,
                            normalized_unit=EmissionRecord.NormalizedUnit.PASSENGER_KM,
                            scope="3",
                            category="flights",
                            activity_value=Decimal('0.0000'),
                            activity_unit="passenger-km",
                            normalized_value=normalized_value,
                            normalized_value_unit="kgCO2e",
                            emission_factor=emission_factor,
                            emission_factor_source=emission_factor_source,
                            period_start=t_date,
                            period_end=t_date + timedelta(hours=3),
                            source_row_id=hashlib_row_id(booking),
                            raw_data=booking,
                            status="REJECTED",
                            workflow_status=EmissionRecord.WorkflowStatus.REJECTED,
                            flag_reason="Navan Ingestion: Booking was CANCELED prior to ingestion.",
                            resolved_facility_id=origin,
                            resolved_facility_country=country_map.get(origin.upper(), 'US') if origin else 'US',
                            is_suspicious=True,
                            validation_errors=["Navan Ingestion: Booking was CANCELED prior to ingestion."],
                            deduplication_key=dedup_key
                        )
                        
                        AuditLog.objects.create(
                            activity=activity,
                            action=AuditLog.AuditAction.CREATE,
                            changed_by=raw.ingested_by,
                            performed_by=user_performer,
                            reason="Auto-ingested canceled flight record (excluded from ledger)."
                        )
                        records.append(activity)
                    continue

            res = normalize_flight(None, origin, dest, cabin, coords_map)
            if res.success:
                norm_qty = Decimal(str(res.canonical_qty))
                norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                normalized_value = Decimal(str(res.kgco2e))
                emission_factor = Decimal(str(res.emission_factor))
                emission_factor_source = res.emission_factor_source
                scope_str = str(res.scope)
                category = "flights"
                if res.flag:
                    validation_errors.append(res.flag)
            else:
                norm_qty = Decimal('0.0000')
                norm_unit = EmissionRecord.NormalizedUnit.PASSENGER_KM
                normalized_value = Decimal('0.0000')
                emission_factor = Decimal('0.0000')
                emission_factor_source = ""
                scope_str = "3"
                category = "flights"
                is_suspicious = True
                validation_errors.append(res.flag)

            dedup_key = f"travel_navan_{trip_id}_{idx}"
            status = "FLAGGED" if is_suspicious else "PENDING_REVIEW"
            
            activity = EmissionRecord.objects.create(
                tenant=tenant,
                raw_payload=raw,
                batch=batch,
                source_row_index=idx,
                unique_transaction_id=trip_id,
                scope_category=EmissionRecord.ScopeCategory.SCOPE_3_TRAVEL,
                start_date=t_date,
                end_date=t_date + timedelta(hours=3),
                raw_quantity=Decimal('0.0000'),
                raw_unit='passenger-km',
                normalized_quantity=norm_qty,
                normalized_unit=norm_unit,
                scope=scope_str,
                category=category,
                activity_value=Decimal('0.0000'),
                activity_unit="passenger-km",
                normalized_value=normalized_value,
                normalized_value_unit="kgCO2e",
                emission_factor=emission_factor,
                emission_factor_source=emission_factor_source,
                period_start=t_date,
                period_end=t_date + timedelta(hours=3),
                source_row_id=hashlib_row_id(booking),
                raw_data=booking,
                status=status,
                flag_reason="; ".join(validation_errors),
                resolved_facility_id=origin,
                resolved_facility_country=country_map.get(origin.upper(), 'US') if origin else 'US',
                is_suspicious=is_suspicious,
                validation_errors=validation_errors,
                deduplication_key=dedup_key
            )
            
            AuditLog.objects.create(
                activity=activity,
                action=AuditLog.AuditAction.CREATE,
                changed_by=raw.ingested_by,
                performed_by=user_performer,
                reason="Auto-ingested flight record from Navan TMC payload."
            )
            records.append(activity)
            
        except Exception as e:
            errors.append(f"Navan element {idx} error: {str(e)}")
            
    return len(records), errors

def hashlib_row_id(row_dict: dict) -> str:
    content = json.dumps(row_dict, sort_keys=True, default=str)
    return hashlib_row_id_str(content)

def hashlib_row_id_str(content_str: str) -> str:
    import hashlib
    return hashlib.sha256(content_str.encode('utf-8')).hexdigest()[:16]
